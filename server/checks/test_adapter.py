"""Small conversion and cancellation fixtures; no semantic split pipeline or model network."""
import asyncio
import json
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from src.tasks import CURRENT, TaskManager, model_request, retained_thread
from src.impl.requirements_split import _convert_sync
from src.split.runner import _run_script, SplitRunError
from src.split.runner import _execute_script
import sys


class AdapterTests(unittest.IsolatedAsyncioTestCase):
    async def test_host_model_queue_keeps_exact_messages(self):
        job = {'pending': {}}
        token = CURRENT.set(job)
        try:
            request = asyncio.create_task(model_request('system only', 'small material'))
            await asyncio.sleep(0)
            [(request_id, (body, future))] = job['pending'].items()
            self.assertEqual(body['messages'], [{'role': 'system', 'content': 'system only'}, {'role': 'user', 'content': 'small material'}])
            future.set_result('{"mock":true}')
            self.assertEqual(await request, '{"mock":true}')
            self.assertFalse(job['pending'])
        finally:
            CURRENT.reset(token)

    async def test_cancellation_waits_for_active_thread(self):
        finished = threading.Event()
        def convert():
            time.sleep(0.08)
            finished.set()
        task = asyncio.create_task(retained_thread(convert))
        await asyncio.sleep(0.01)
        task.cancel()
        with self.assertRaises(asyncio.CancelledError):
            await task
        self.assertTrue(finished.is_set())

    async def test_cancel_stops_only_owned_script_and_retains_stdout(self):
        cancelled = threading.Event()
        token = CURRENT.set({'cancel_requested': cancelled})
        try:
            task = asyncio.create_task(asyncio.to_thread(_execute_script,
                [sys.executable, '-u', '-c', 'import time; print("tiny child output", flush=True); time.sleep(10)'], 20))
            await asyncio.sleep(0.25)
            cancelled.set()
            result = await task
            self.assertEqual(result.returncode, 130)
            self.assertIn('tiny child output', result.stdout)
        finally:
            CURRENT.reset(token)

    async def test_small_docx_and_text_use_only_fixed_categories(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for category in ('input', 'intermediate', 'final', 'logs'):
                (root / category).mkdir()
            source = root / 'tiny.docx'
            doc = Document(); doc.add_heading('订单', 1); doc.add_paragraph('用户查询订单。'); doc.save(source)
            result = _convert_sync({'path': str(source), '_run_directory': str(root)})
            self.assertIn('用户查询订单', Path(result['markdownPath']).read_text(encoding='utf-8'))
            self.assertTrue((root / 'input' / source.name).exists())
            self.assertEqual(Path(result['outputDir']), root / 'intermediate')
            _convert_sync({'text': '# 1 订单\n用户查询订单。', '_run_directory': str(root)})
            self.assertTrue((root / 'input' / 'material.txt').exists())
            self.assertFalse(list((root / 'final').iterdir()))

    async def test_script_failure_keeps_full_stdout_and_stage_snapshot(self):
        import subprocess
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); (root / 'final').mkdir(); (root / 'final' / 'tiny.md').write_text('retained')
            token = CURRENT.set({'directory': root})
            try:
                with patch('src.split.runner._execute_script', return_value=subprocess.CompletedProcess([], 2, 'all stdout', 'all stderr')):
                    with self.assertRaises(SplitRunError):
                        _run_script('validate_chunks.py', root, report=[], timeout=1)
                [snapshot] = (root / 'intermediate' / 'stages').iterdir()
                self.assertEqual(json.loads((snapshot / 'execution.json').read_text())['stdout'], 'all stdout')
                self.assertEqual((snapshot / 'final' / 'tiny.md').read_text(), 'retained')
            finally:
                CURRENT.reset(token)


if __name__ == '__main__':
    unittest.main()
