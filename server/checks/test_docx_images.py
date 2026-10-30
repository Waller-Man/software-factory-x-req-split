"""Real DOCX image conversion fixtures; no model calls or requirement pipeline."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from src.convert.base import convert_file
from src.convert.docx import VML_IMAGEDATA


def image_bytes(color):
    output = BytesIO()
    Image.new('RGB', (2, 2), color).save(output, format='PNG')
    return output.getvalue()


def image_relation(document, blob):
    rid, _ = document.part.get_or_add_image(BytesIO(blob))
    return rid


def add_vml(paragraph, rid):
    pict = OxmlElement('w:pict')
    shape = etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}shape')
    image = etree.SubElement(shape, VML_IMAGEDATA)
    if rid is not None:
        image.set(qn('r:id'), rid)
    paragraph.add_run()._r.append(pict)


class DocxImageTests(unittest.TestCase):
    def convert(self, document, expected_blobs, *, table=False):
        with tempfile.TemporaryDirectory(prefix='req-docx-images-') as folder:
            root = Path(folder)
            source = root / 'tiny.docx'
            document.save(source)
            output = root / 'intermediate'
            manifest = convert_file(source, output)
            markdown = (output / 'document.md').read_text(encoding='utf-8')
            self.assertIn('订单正文', markdown)
            self.assertEqual(len(manifest['assets']), len(expected_blobs))
            for asset, expected in zip(manifest['assets'], expected_blobs):
                path = output / asset['path']
                self.assertEqual(path.read_bytes(), expected)
                self.assertIn(asset['path'], markdown)
                self.assertEqual(asset['section_id'], '1')
            if table:
                self.assertIn('<table>', markdown)
                self.assertIn('<img src=', markdown)
            return manifest

    def document(self):
        document = Document()
        document.add_heading('1 订单', level=1)
        return document

    def test_vml_only_paragraph_uses_its_image_relation(self):
        document = self.document()
        blob = image_bytes('red')
        add_vml(document.add_paragraph('订单正文'), image_relation(document, blob))
        self.assertFalse(self.convert(document, [blob])['warnings'])

    def test_mixed_drawingml_and_vml_do_not_reuse_previous_relation(self):
        document = self.document()
        first, second = image_bytes('red'), image_bytes('blue')
        paragraph = document.add_paragraph('订单正文')
        paragraph.add_run().add_picture(BytesIO(first))
        add_vml(paragraph, image_relation(document, second))
        self.assertFalse(self.convert(document, [first, second])['warnings'])

    def test_multiple_vml_images_each_use_their_own_relation(self):
        document = self.document()
        first, second = image_bytes('red'), image_bytes('green')
        paragraph = document.add_paragraph('订单正文')
        add_vml(paragraph, image_relation(document, first))
        add_vml(paragraph, image_relation(document, second))
        self.assertFalse(self.convert(document, [first, second])['warnings'])

    def test_vml_without_relation_does_not_duplicate_drawingml_image(self):
        document = self.document()
        blob = image_bytes('red')
        paragraph = document.add_paragraph('订单正文')
        paragraph.add_run().add_picture(BytesIO(blob))
        add_vml(paragraph, None)
        self.convert(document, [blob])

    def test_unresolved_vml_relation_keeps_text_and_reports_warning(self):
        document = self.document()
        add_vml(document.add_paragraph('订单正文'), 'rIdMissing')
        manifest = self.convert(document, [])
        self.assertTrue(any('rIdMissing' in warning for warning in manifest['warnings']))

    def test_vml_image_in_table_is_extracted(self):
        document = self.document()
        blob = image_bytes('blue')
        paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        paragraph.add_run('订单正文')
        add_vml(paragraph, image_relation(document, blob))
        self.assertFalse(self.convert(document, [blob], table=True)['warnings'])


if __name__ == '__main__':
    unittest.main()
